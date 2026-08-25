from openpyxl import Workbook
from docx import Document
from enterprise_data_context.ingestion.excel import ExcelParser
from enterprise_data_context.ingestion.word import WordParser

def test_excel_regions_and_two_headers(tmp_path):
    p=tmp_path/"x.xlsx"; wb=Workbook(); ws=wb.active
    ws.append(["Model","","Layer"]); ws.append(["Model Name","Description","Layer"])
    ws.append(["LTE MR","MR detail","ODS"]); ws.append([])
    ws.append(["Metric","Formula"]); ws.append(["RSRP","AVG(x)"]); wb.save(p)
    ir=ExcelParser().parse("x",p,"asset_catalog")
    assert len(ir.elements)==2
    assert ir.elements[0].metadata["header_rows"] in (1,2)

def test_word_heading_table_binding(tmp_path):
    p=tmp_path/"x.docx"; d=Document()
    d.add_heading("Coverage",1); d.add_paragraph("Description")
    t=d.add_table(rows=2,cols=2); t.cell(0,0).text="指标"; t.cell(0,1).text="公式"
    t.cell(1,0).text="RSRP"; t.cell(1,1).text="AVG(x)"; d.save(p)
    ir=WordParser().parse("x",p,"kpi_definition")
    table=[e for e in ir.elements if e.type=="table"][0]
    assert "Coverage" in (table.source.section or "")
