from openpyxl import Workbook
from enterprise_data_context.compiler import ContextCompiler
from enterprise_data_context.runtime import from_compiled
from explore_agent import ExploreAgent

def test_e2e(tmp_path):
    k=tmp_path/"kpi.xlsx"; wb=Workbook(); ws=wb.active
    ws.append(["指标名称","指标描述","公式","维度","数据源"])
    ws.append(["RSRP","无线信号强度","AVG(RSRP)","Cell,Grid","LTE Periodic MR"]); wb.save(k)

    m=tmp_path/"models.xlsx"; wb=Workbook(); ws=wb.active
    ws.append(["模型名称","模型描述","分层","主题","粒度"])
    ws.append(["LTE Periodic MR","LTE MR","ODS","Wireless Coverage","Subscriber,Cell,Time"]); wb.save(m)

    compiled=ContextCompiler().compile([
        {"id":"k","path":str(k),"type":"kpi_definition"},
        {"id":"m","path":str(m),"type":"asset_catalog"},
    ])
    runtime=from_compiled(compiled)
    bundle=ExploreAgent(runtime.retrieval).explore("RSRP 有哪些现有模型可以提供？")
    assert any(x["name"]=="RSRP" for x in bundle.primary_contexts)
    assert any(x["name"]=="LTE Periodic MR" for x in bundle.primary_contexts)
    assert bundle.environment["binding_required"] is True
    assert bundle.environment["availability_state"] == "UNSUPPORTED"
    assert "environment_availability" in bundle.missing_context

def test_partial_knowledge_is_explicit(tmp_path):
    k=tmp_path/"kpi.xlsx"; wb=Workbook(); ws=wb.active
    ws.append(["指标名称","指标描述"]); ws.append(["RSRP","无线信号强度"]); wb.save(k)
    compiled=ContextCompiler().compile([{"id":"k","path":str(k),"type":"kpi_definition"}])
    bundle=ExploreAgent(from_compiled(compiled).retrieval).explore("RSRP 有哪些现有模型可以提供？")
    assert "models" in bundle.missing_context
