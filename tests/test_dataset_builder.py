from copy import deepcopy
import json
from pathlib import Path
import shutil
import pytest
import yaml
from openpyxl import load_workbook

from evaluation.dataset_builder.fixtures import create_sample
from evaluation.dataset_builder.contracts import load_profiles, Profiles, assert_draft
from evaluation.dataset_builder.builder import build
from evaluation.dataset_builder.diagnostics import load_dataset,validate,case_sources,inspect_entity,inspect_relation,inspect_case
from evaluation.dataset_builder.cli import main
from evaluation.dataset_builder.sampler import sample
from evaluation.dataset_builder.entity_graph import EntityGraph
from evaluation.dataset_builder.source_inventory import inventory
from evaluation.benchmark.preflight import guard_workspace,PreflightError
from evaluation.cases.schema import validate_cases
from evaluation.normalization.entity_registry import EntityRegistry


@pytest.fixture
def dataset(tmp_path):
    raw,parser,profile=create_sample(tmp_path/"input")
    out=tmp_path/"dataset"
    report=build(raw,parser,out,load_profiles(profile))
    return raw,parser,profile,out,report


def test_sample_build_multidocument_motifs_blind_spot_and_draft_only(dataset):
    raw,parser,profile,out,report=dataset; data=load_dataset(out)
    assert set(p.name for p in out.iterdir())=={"cases.draft.yaml","aliases.draft.yaml","candidate_entities.json",
        "candidate_relations.json","source_entity_map.json","review_queue.json","review_queue.xlsx","generation_report.json"}
    assert {c["category"] for c in data["cases"]}=={"metric_to_model","purpose_to_data","model_to_analysis","model_to_business","field_discovery","lineage_impact","negative"}
    assert report["selected"]["parser_blind_spot_cases"]>=1
    assert "metric:metricB" in report["raw_vs_parser"]["raw_independent_only"]
    assert "metric:metricA" in report["raw_vs_parser"]["both"]
    assert report["selected"]["quota_warnings"]
    assert validate(out)["draft_only"]
    assert all(row["review_status"]=="DRAFT" for row in [*data["entities"],*data["relations"],*data["cases"],*data["aliases"]["entities"]])
    assert all("gold" not in case for case in data["cases"])
    purpose=next(c for c in data["cases"] if c["category"]=="purpose_to_data")
    assert purpose["source_span"]=="cross_document" and len(purpose["source_documents"])==2
    assert purpose["cross_document_rationale"] and purpose["difficulty_suggested"]=="hard"
    assert all(row["valid"] for row in case_sources(data,raw)["checks"])


def test_alias_collision_does_not_merge_entities_or_create_identity(dataset):
    data=load_dataset(dataset[3]); collision=next(c for c in data["aliases"]["collisions"] if c["alias"]=="ne")
    assert len(collision["entity_ids"])==2
    assert set(collision["entity_ids"])=={"business-object:objectA","business-object:objectB"}
    graph=EntityGraph(Profiles(),["a.docx"])
    graph.add_entities([{"type":"metrics","name":"SAME","source_documents":["a.docx"],"parser_location":{"parser_file":"x.json","pointer":"/0"}},
                        {"type":"metrics","name":"SAME","source_documents":["a.docx"],"parser_location":{"parser_file":"x.json","pointer":"/1"}}],"parser")
    assert len(graph.entities)==2


def test_parser_supported_is_not_raw_verified_or_approved(tmp_path):
    raw,parser,profile=create_sample(tmp_path/"input")
    path=parser/"sample.json"; value=json.loads(path.read_text())
    value["entities"].append({"id":"invented","name":"not in raw","type":"metrics","review_status":"APPROVED", "source_documents":["models.docx"]})
    path.write_text(json.dumps(value))
    build(raw,parser,tmp_path/"out",load_profiles(profile))
    entity=next(e for e in load_dataset(tmp_path/"out")["entities"] if e["id"]=="metric:invented")
    assert entity["parser_supported"] and not entity["raw_verified"] and entity["review_status"]=="DRAFT"


def test_negative_only_comes_from_explicit_seed_and_remains_unconfirmed(tmp_path,dataset):
    data=load_dataset(dataset[3]); negative=next(c for c in data["cases"] if c["category"]=="negative")
    assert negative["candidate_kind"]=="negative_candidate" and negative["requires_human_confirmation"]
    assert negative["gold_candidate"]==[] and not negative["raw_verified"]
    profile=load_profiles(dataset[2]); profile.generation.negative_candidates=[]
    report=build(dataset[0],dataset[1],tmp_path/"no-negative",profile)
    assert report["all_candidates"]["category_distribution"]["negative"]==0


def test_query_only_paraphrase_cannot_change_answer_candidates(tmp_path,dataset):
    received=[]
    def paraphrase(query):
        assert isinstance(query,str); received.append(query)
        return "请依据资料回答："+query
    out=tmp_path/"rewritten"
    build(dataset[0],dataset[1],out,load_profiles(dataset[2]),paraphraser=paraphrase)
    previous={c["case_id"]:c for c in load_dataset(dataset[3])["cases"]}
    current=load_dataset(out)["cases"]
    assert received
    for case in current:
        assert case["gold_candidate"]==previous[case["case_id"]]["gold_candidate"]
        assert case["support_relation_candidates"]==previous[case["case_id"]]["support_relation_candidates"]


def test_sampler_caps_duplicates_and_motif_diversity(dataset):
    rows=load_dataset(dataset[3])["cases"]
    profile=load_profiles(dataset[2]).generation
    profile.max_cases_per_document=1
    selected,report=sample(rows,profile)
    assert selected and len(selected)<=2 and any(e["reason"]=="document_cap" for e in report["excluded"])
    profile.max_cases_per_document=100; profile.max_cases_per_entity=100; profile.max_cases_per_topic=100
    profile.max_cases_per_motif=1
    selected,report=sample(rows,profile)
    assert len({c["motif_signature"] for c in selected})==len(selected)
    clone=deepcopy(rows[0]); clone["case_id"]="distinct-id-same-query"
    selected,report=sample([rows[0],clone],profile)
    assert len(selected)==1 and report["excluded"]


def test_draft_enforcement_and_gold_leakage_guard(dataset,tmp_path):
    data=load_dataset(dataset[3]); promoted=deepcopy(data["cases"]); promoted[0]["review_status"]="APPROVED"
    with pytest.raises(ValueError,match="DRAFT"): assert_draft(promoted)
    with pytest.raises(ValueError,match="never approved"): assert_draft({"gold":{}})
    with pytest.raises(ValueError,match="approved"):
        validate_cases(data["cases"],EntityRegistry(data["aliases"]["entities"]),{"models.docx","dictionary.xlsx"})
    workspace=tmp_path/"native"; workspace.mkdir()
    shutil.copyfile(dataset[3]/"aliases.draft.yaml",workspace/"cache.txt")
    with pytest.raises(PreflightError): guard_workspace(workspace,[dataset[3]/"cases.draft.yaml",dataset[3]/"aliases.draft.yaml"])


def test_cli_diagnostics_and_inspectors(dataset,capsys):
    raw,parser,profile,out,_=dataset; data=load_dataset(out)
    commands=[ ["inventory","--raw-dir",str(raw)], ["inspect_raw_structure","--raw-dir",str(raw)],
        ["inspect_parser_json","--parser-json-dir",str(parser),"--profile",str(profile)],
        *[[cmd,"--dataset-dir",str(out)] for cmd in ("validate","report","compare_raw_vs_parser")],
        ["validate_case_sources","--dataset-dir",str(out),"--raw-dir",str(raw)],
        ["inspect_entity","--dataset-dir",str(out),"--id",data["entities"][0]["id"]],
        ["inspect_relation","--dataset-dir",str(out),"--id",data["relations"][0]["id"]],
        ["inspect_case","--dataset-dir",str(out),"--id",data["cases"][0]["case_id"]]]
    for command in commands:
        assert main(command)==0
        assert isinstance(json.loads(capsys.readouterr().out),dict)


def test_profiles_adapt_custom_parser_envelope_without_core_changes(tmp_path,dataset):
    parser=tmp_path/"adapted-parser"; parser.mkdir()
    value=json.loads((dataset[1]/"sample.json").read_text())
    (parser/"different.json").write_text(json.dumps({"knowledge":{"objects":value["entities"],"links":value["relations"]}}))
    profile=load_profiles(dataset[2]); profile.parser.entities[0]["records"]="/knowledge/objects/*"; profile.parser.relations[0]["records"]="/knowledge/links/*"
    report=build(dataset[0],parser,tmp_path/"adapted",profile)
    assert report["all_candidates"]["candidate_entities"]==dataset[4]["all_candidates"]["candidate_entities"]


def test_workbook_is_review_only_and_text_is_not_a_formula(dataset):
    workbook=load_workbook(dataset[3]/"review_queue.xlsx")
    assert workbook.sheetnames==["Cases","Entities","Relations","Aliases"]
    for sheet in workbook:
        assert sheet.freeze_panes=="A2"
        for row in sheet.iter_rows(min_row=2):
            assert row[8].value=="DRAFT" and row[9].value is None
            assert all(cell.data_type!="f" for cell in row)
    workbook.close()


def test_excel_review_export_does_not_execute_formula_names(tmp_path):
    from evaluation.dataset_builder.review_export import write_review
    entity={"id":"metric:unsafe-text","name":"=1+1","type":"metrics","source_documents":[],"source_locations":[]}
    write_review(tmp_path,[],[entity],[],{"entities":[]})
    book=load_workbook(tmp_path/"review_queue.xlsx")
    assert book["Entities"]["C2"].value=="=1+1" and book["Entities"]["C2"].data_type=="s"
    book.close()


def test_formal_manifest_distribution_retains_parser_origin_counts(dataset):
    from evaluation.cases.schema import distribution
    rows=load_dataset(dataset[3])["cases"]
    actual=distribution(rows)
    assert actual["parser_supported_cases"]==sum(r["parser_supported"] for r in rows)
    assert actual["parser_blind_spot_cases"]>=1 and actual["parser_provenance_unknown_cases"]==0


def test_invalid_relation_types_and_unsupported_raw_are_explicit(tmp_path,dataset):
    from evaluation.dataset_builder.source_inventory import inspect_raw_structure
    raw=tmp_path/"legacy-raw"; raw.mkdir(); (raw/"legacy.doc").write_bytes(b"synthetic-unsupported")
    result=inspect_raw_structure(raw)
    assert result["units"]==[] and result["diagnostics"][0]["status"]=="UNSUPPORTED"
    parser=tmp_path/"parser-copy"; shutil.copytree(dataset[1],parser)
    path=parser/"sample.json"; value=json.loads(path.read_text())
    value["relations"].append({"source":{"type":"physical_models","id":"modelA"},"relation":"supported_by",
                                "target":{"type":"metrics","id":"metricA"},"source_documents":["models.docx"]})
    path.write_text(json.dumps(value))
    result=build(dataset[0],parser,tmp_path/"invalid-rel-out",load_profiles(dataset[2]))
    assert any(r.get("code")=="relation_type_mismatch" for r in result["diagnostics"])


def test_unhandled_source_location_precision_cannot_claim_raw_verification():
    from evaluation.dataset_builder.raw_verifier import locator_matches
    unit={"location":{"document":"x.xlsx","kind":"excel_row","sheet":"S","row":2}}
    assert locator_matches({"document":"x.xlsx","row":2},unit)
    assert not locator_matches({"document":"x.xlsx","row":2,"cell":"Z2"},unit)


def test_inputs_immutable_output_isolated_and_empty_required(dataset,tmp_path):
    raw,parser,profile,out,report=dataset
    assert inventory(raw)==report["raw_inventory"] and inventory(parser,True)==report["parser_inventory"]
    from evaluation.benchmark.raw_corpus import fingerprint
    assert fingerprint(report["raw_inventory"])==report["raw_hash"]
    with pytest.raises(ValueError,match="outside"): build(raw,parser,raw/"review",load_profiles(profile))
    with pytest.raises(ValueError,match="empty"): build(raw,parser,out,load_profiles(profile))
    (raw/"symlink.docx").symlink_to(raw/"models.docx")
    with pytest.raises(ValueError,match="symlinks"): inventory(raw)
