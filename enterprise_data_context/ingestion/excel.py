from pathlib import Path
from openpyxl import load_workbook
from .base import Parser
from enterprise_data_context.models import DocumentIR, DocumentElement, SourceLocation

def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())

def row_blank(row):
    return all(blank(v) for v in row)

def trim(row):
    row = list(row)
    while row and blank(row[-1]):
        row.pop()
    return row

def detect_regions(ws):
    """Conservative rectangular regions split by fully blank rows."""
    rows = [trim(r) for r in ws.iter_rows(values_only=True)]
    out, start = [], None
    for i, row in enumerate(rows, 1):
        if not row_blank(row) and start is None:
            start = i
        elif row_blank(row) and start is not None:
            out.append((start, i - 1))
            start = None
    if start is not None:
        out.append((start, len(rows)))
    return out

def infer_header_rows(data):
    if not data:
        return 0
    # Handle one or two header rows conservatively.
    first = [x for x in data[0] if not blank(x)]
    if len(first) < 2:
        return 0
    str_ratio = sum(isinstance(x, str) for x in first) / max(1, len(first))
    if str_ratio < 0.6:
        return 0
    if len(data) > 1:
        second = [x for x in data[1] if not blank(x)]
        second_ratio = sum(isinstance(x, str) for x in second) / max(1, len(second)) if second else 0
        # Merged/meta style first row + descriptive second row.
        if len(first) <= max(2, len(second)//2) and second_ratio >= 0.6:
            return 2
    return 1

class ExcelParser(Parser):
    extensions = (".xlsx", ".xlsm")

    def parse(self, source_id, path, source_type="unknown"):
        path = Path(path)
        wb = load_workbook(path, data_only=False, read_only=False)
        elems = []
        for ws in wb.worksheets:
            for n, (start, end) in enumerate(detect_regions(ws), 1):
                data = [trim([ws.cell(r, c).value for c in range(1, ws.max_column+1)]) for r in range(start, end+1)]
                if not data:
                    continue
                hcount = infer_header_rows(data)
                headers = []
                if hcount == 1:
                    headers = [str(x).strip() if x is not None else "" for x in data[0]]
                elif hcount == 2:
                    width = max(len(data[0]), len(data[1]))
                    top = data[0] + [""] * (width-len(data[0]))
                    bot = data[1] + [""] * (width-len(data[1]))
                    headers = [
                        " / ".join([str(x).strip() for x in (top[i], bot[i]) if not blank(x)])
                        for i in range(width)
                    ]
                body = data[hcount:]
                elems.append(DocumentElement(
                    type="table", id=f"{ws.title}:region-{n}", title=f"{ws.title} region {n}",
                    headers=headers, rows=body,
                    metadata={
                        "sheet": ws.title, "start_row": start, "end_row": end,
                        "header_rows": hcount,
                        "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
                        "hidden_rows": [i for i, d in ws.row_dimensions.items() if d.hidden],
                        "hidden_columns": [i for i, d in ws.column_dimensions.items() if d.hidden],
                    },
                    source=SourceLocation(source_id, str(path), sheet=ws.title, table=f"region-{n}", row=start)
                ))
        return DocumentIR(source_id, str(path), source_type, {"format": path.suffix.lower()}, elems)
