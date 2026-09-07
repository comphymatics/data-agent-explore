"""Deterministic parser for the documented synthetic table layout ONLY.

Not the production Domain Parser. This executable proves the raw -> Template ->
unchanged Compiler -> Explore path without substituting precomputed contexts.
"""
import json
from pathlib import Path
import sys
from docx import Document
from openpyxl import load_workbook
from evaluation.benchmark.raw_corpus import inventory

VERSION="synthetic-table-parser/v1"


def parse(root,output):
    models={}; columns=[]; trace=[]
    for file in inventory(root):
        path=Path(root)/file["path"]
        if path.suffix==".docx":
            doc=Document(path)
            for table in doc.tables:
                header=[c.text for c in table.rows[0].cells]
                if header!=["table_name","table_description","layer","logic_model_name"]:
                    raise ValueError("unsupported synthetic Word table layout")
                for row in table.rows[1:]:
                    value=dict(zip(header,[c.text for c in row.cells])); models[value["table_name"]]=value
        elif path.suffix==".xlsx":
            workbook=load_workbook(path,read_only=True,data_only=True)
            for sheet in workbook:
                rows=iter(sheet.values); header=next(rows)
                if tuple(header)!=("table_name","column_name","column_description","data_type","catagory","corresponding_counter_or_dim"):
                    raise ValueError("unsupported synthetic Excel layout")
                columns += [dict(zip(header,[v if v is not None else "" for v in row])) for row in rows]
            workbook.close()
        else: raise ValueError("synthetic parser supports only its .docx/.xlsx layout")
        trace.append({"phase":"deterministic_parse","file":file,"llm_tokens":0})
    for col in columns:
        if col["table_name"] not in models: raise ValueError("dictionary references an undeclared model")
    tables=[]
    for name,model in models.items():
        fields=[]
        for col in columns:
            if col["table_name"]==name:
                fields.append({**{k:v for k,v in col.items() if k!="table_name"},"content_decription":col["column_description"],
                    "unit":"","source_columns":[],"processing_logic":"","supported_scenarios":""})
        tables.append({**model,"model_type":"","granularity":"","storage":"","domain":"","topic":"","app_name":"",
                       "source":{"source_tables":[],"processing_logic":""},"columns":fields})
    output=Path(output); output.mkdir(parents=True,exist_ok=True)
    (output/"tables.json").write_text(json.dumps({"tables":tables},ensure_ascii=False,indent=2))
    return {"protocol":"raw-e2e-driver/v1","status":"OK","parser_version":VERSION,"consumed_files":inventory(root),
            "usage_complete":True,"llm_calls":[],"trace":trace,"coverage":"PARTIAL; synthetic table layout only; narrative semantics not extracted"}


if __name__=="__main__":
    request=json.load(sys.stdin)
    if request.get("phase")!="parse": raise ValueError("parse phase required")
    print(json.dumps(parse(request["corpus_path"],request["output_dir"])))
