"""Create original synthetic Word/Excel files; no Template/Gold in the raw directory."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook


def create_raw(path):
    root=Path(path); root.mkdir(parents=True,exist_ok=False)
    doc=Document()
    doc.add_heading("Synthetic LTE model handbook",0)
    table=doc.add_table(rows=1,cols=4)
    headers=["table_name","table_description","layer","logic_model_name"]
    for cell,text in zip(table.rows[0].cells,headers): cell.text=text
    for values in [("LTE_PERIODIC_MR","LTE小区测量数据，支持地铁弱覆盖分析","ODS","LTE_MR_LOGICAL"),
                   ("CELL_HOURLY","LTE_PERIODIC_MR 的下游小时汇总模型","ADS","")]:
        for cell,text in zip(table.add_row().cells,values): cell.text=text
    doc.add_paragraph("业务对象 Cell 对应 LTE_PERIODIC_MR。修改 LTE_PERIODIC_MR.CELL_ID 会影响 CELL_HOURLY。")
    doc.add_paragraph("地铁弱覆盖分析使用 RSRP 指标以及 LTE_PERIODIC_MR 的 CELL_ID 字段。")
    doc.save(root/"models.docx")
    wb=Workbook(); ws=wb.active; ws.title="Columns"
    ws.append(["table_name","column_name","column_description","data_type","catagory","corresponding_counter_or_dim"])
    ws.append(["LTE_PERIODIC_MR","CELL_ID","小区标识","STRING","dimension","Cell"])
    ws.append(["LTE_PERIODIC_MR","RSRP_VALUE","接收信号功率","DOUBLE","metric","RSRP"])
    ws.append(["CELL_HOURLY","CELL_ID","小区标识","STRING","dimension","Cell"])
    wb.save(root/"dictionary.xlsx")
    return root
