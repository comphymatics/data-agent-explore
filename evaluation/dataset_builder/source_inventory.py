"""Read-only, deterministic structural inspection, not a per-document semantic parser."""
import hashlib
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from .contracts import fingerprint


def inventory(root, parser=False):
    root=Path(root).resolve()
    if not root.is_dir(): raise ValueError("input directory does not exist")
    rows=[]
    allowed={".json"} if parser else {".doc", ".docx", ".xls", ".xlsx"}
    for path in sorted(root.rglob("*")):
        if path.is_symlink(): raise ValueError("input symlinks are forbidden")
        if not path.is_file(): continue
        if path.suffix.lower() not in allowed: raise ValueError("unexpected input format: "+path.name)
        rows.append({"path":path.relative_to(root).as_posix(),"bytes":path.stat().st_size,
                     "sha256":hashlib.sha256(path.read_bytes()).hexdigest()})
    if not rows: raise ValueError("empty input directory")
    return rows


def read_raw(root, files):
    units=[]; diagnostics=[]
    def unit(document,kind,location,text,record=None):
        units.append({"document":document,"kind":kind,"location":{"document":document,"kind":kind,**location},
                      "text":text,"record":record or {}})
    for file in files:
        name=file["path"]; path=Path(root)/name
        try:
            if path.suffix.lower()==".docx":
                document=Document(path)
                for i,paragraph in enumerate(document.paragraphs,1):
                    if paragraph.text.strip(): unit(name,"word_paragraph",{"paragraph":i},paragraph.text)
                for ti,table in enumerate(document.tables,1):
                    for ri,row in enumerate(table.rows,1):
                        values=[cell.text for cell in row.cells]
                        unit(name,"word_table_row",{"table":ti,"row":ri}," | ".join(values),
                             {str(i):value for i,value in enumerate(values,1)})
                diagnostics.append({"document":name,"status":"STRUCTURAL_ONLY","coverage":"PARTIAL",
                                    "scope":"top-level paragraphs and tables; headers, drawings and other structures need an adapter"})
            elif path.suffix.lower()==".xlsx":
                workbook=load_workbook(path,read_only=True,data_only=False)
                try:
                    for sheet in workbook:
                        for ri,row in enumerate(sheet.iter_rows(values_only=True),1):
                            values=[str(v) if v is not None else "" for v in row]
                            if any(values): unit(name,"excel_row",{"sheet":sheet.title,"row":ri}," | ".join(values),
                                                {str(i):value for i,value in enumerate(values,1)})
                finally: workbook.close()
                diagnostics.append({"document":name,"status":"STRUCTURAL_ONLY","coverage":"PARTIAL",
                                    "scope":"worksheet cell values/formulas; comments, drawings and semantic relations need an adapter"})
            else:
                diagnostics.append({"document":name,"status":"UNSUPPORTED","reason":"inject a RawReader for legacy .doc/.xls"})
        except Exception as exc:
            diagnostics.append({"document":name,"status":"UNREADABLE","reason":type(exc).__name__})
    for unit_row in units:
        unit_row["unit_id"]="unit-"+fingerprint(unit_row["location"])[:20]
    return units,diagnostics


def inspect_raw_structure(root, reader=read_raw):
    files=inventory(root); units,diagnostics=reader(Path(root),files)
    return {"files":files,"corpus_hash":fingerprint(files),"units":units,"diagnostics":diagnostics}
