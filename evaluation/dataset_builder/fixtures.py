"""Synthetic/sample-only material generator. Never used by production parsing."""
import json
from pathlib import Path
from docx import Document
from openpyxl import Workbook
import yaml


def create_sample(directory):
    root=Path(directory); raw=root/"raw"; parser=root/"parser-json"
    raw.mkdir(parents=True); parser.mkdir()
    word_entities=[
        ["purposes","purposeA","Coverage Review","","LTE","coverage"],
        ["physical_models","modelA","RADIO_SAMPLE","","LTE","coverage"],
        ["physical_models","modelB","DAILY_SAMPLE","","LTE","aggregation"],
        ["business_objects","objectA","Cell","NE","LTE","coverage"],
        ["business_objects","objectB","Element","NE","LTE","inventory"],
        ["logical_models","logicA","Radio observation","","LTE","coverage"],
    ]
    excel_entities=[
        ["metrics","metricA","Receive Power","POWER","LTE","coverage"],
        ["fields","fieldA","RADIO_SAMPLE.cell_id","","LTE","coverage"],
        ["fields","fieldB","RADIO_SAMPLE.power_value","","LTE","coverage"],
        ["metrics","metricB","Signal Variability","VARIABILITY","LTE","aggregation"],
    ]
    word_relations=[
        ["purposes","purposeA","requires_metric","metrics","metricA"],
        ["physical_models","modelA","belongs_to_object","business_objects","objectA"],
        ["physical_models","modelA","downstream","physical_models","modelB"],
        ["logical_models","logicA","implemented_by","physical_models","modelA"],
    ]
    excel_relations=[
        ["metrics","metricA","supported_by","physical_models","modelA"],
        ["physical_models","modelA","has_field","fields","fieldA"],
        ["physical_models","modelA","has_field","fields","fieldB"],
        ["metrics","metricB","supported_by","physical_models","modelB"],
    ]
    entity_header=["type","id","name","alias","technology","topic"]
    relation_header=["source_type","source_id","relation","target_type","target_id"]
    doc=Document(); doc.add_paragraph("Synthetic source only; not enterprise documentation.")
    for header,rows in ((entity_header,word_entities),(relation_header,word_relations)):
        table=doc.add_table(rows=1,cols=len(header))
        for cell,value in zip(table.rows[0].cells,header): cell.text=value
        for values in rows:
            for cell,value in zip(table.add_row().cells,values): cell.text=value
    doc.save(raw/"models.docx")
    book=Workbook(); book.remove(book.active)
    for title,header,rows in (("Entities",entity_header,excel_entities),("Relations",relation_header,excel_relations)):
        sheet=book.create_sheet(title); sheet.append(header)
        for row in rows: sheet.append(row)
    book.save(raw/"dictionary.xlsx")
    entities=[]; relations=[]
    for filename,kind,entity_rows,relation_rows in (("models.docx","word_table_row",word_entities,word_relations),
                                                   ("dictionary.xlsx","excel_row",excel_entities,excel_relations)):
        for i,row in enumerate(entity_rows,2):
            if row[1]=="metricB": continue  # Deliberate Parser blind spot, present in independently readable raw.
            location={"document":filename,"kind":kind,"row":i,**({"table":1} if kind=="word_table_row" else {"sheet":"Entities"})}
            entities.append({"type":row[0],"id":row[1],"name":row[2],"aliases":[row[3]] if row[3] else [],
                             "facets":{"technology":row[4],"topic":row[5]},"source_documents":[filename],"source_locations":[location]})
        for i,row in enumerate(relation_rows,2):
            if row[1]=="metricB": continue
            location={"document":filename,"kind":kind,"row":i,**({"table":2} if kind=="word_table_row" else {"sheet":"Relations"})}
            relations.append({"source":{"type":row[0],"id":row[1]},"relation":row[2],"target":{"type":row[3],"id":row[4]},
                              "source_documents":[filename],"source_locations":[location]})
    (parser/"sample.json").write_text(json.dumps({"entities":entities,"relations":relations},ensure_ascii=False,indent=2))
    entity_mapping={"type":"/1","id":"/2","name":"/3","aliases":"/4","facets":{"technology":"/5","topic":"/6"}}
    relation_mapping={"source":{"type":"/1","id":"/2"},"relation":"/3","target":{"type":"/4","id":"/5"}}
    profile={"source":{"raw_entities":[
        {"kind":"word_table_row","file_glob":"*.docx","location_match":{"table":1},"start_row":2,"mapping":entity_mapping},
        {"kind":"excel_row","file_glob":"*.xlsx","sheet_glob":"Entities","start_row":2,"mapping":entity_mapping}],
        "raw_relations":[
        {"kind":"word_table_row","file_glob":"*.docx","location_match":{"table":2},"start_row":2,"mapping":relation_mapping},
        {"kind":"excel_row","file_glob":"*.xlsx","sheet_glob":"Relations","start_row":2,"mapping":relation_mapping}]},
        "generation":{"negative_candidates":[{"query":"当前资料能否确认存在 VoNR 掉话分析模型？",
            "source_documents":["models.docx","dictionary.xlsx"],"review_instruction":"Test-only review candidate; manually confirm the material-scoped absence."}]}}
    path=root/"sample-profile.yaml"; path.write_text(yaml.safe_dump(profile,sort_keys=False,allow_unicode=True))
    return raw,parser,path
