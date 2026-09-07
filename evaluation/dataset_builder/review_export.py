"""Review packets preserve full JSON; Excel is a convenient view, not an approval importer."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill


def write_review(directory,cases,entities,relations,aliases):
    directory=Path(directory)
    packet={"review_status":"DRAFT","instructions":"Verify raw sources, semantic identity, relations and query meaning. Builder never approves.",
            "cases":cases,"entities":entities,"relations":relations,"aliases":aliases}
    (directory/"review_queue.json").write_text(json.dumps(packet,ensure_ascii=False,indent=2))
    workbook=Workbook(); workbook.remove(workbook.active)
    sheets={"Cases":cases,"Entities":entities,"Relations":relations,"Aliases":aliases["entities"]}
    headers=["id","category_or_type","name_or_query","origin","parser_supported","raw_verified","source_documents",
             "source_locations","review_status","review_decision","reviewer","review_notes"]
    for title,records in sheets.items():
        sheet=workbook.create_sheet(title); sheet.append(headers)
        for record in records:
            values=[record.get("case_id",record.get("id",record.get("canonical_id"))),record.get("category",record.get("type",record.get("relation"))),
                    record.get("query",record.get("name","")),record.get("candidate_origin"),record.get("parser_supported"),
                    record.get("raw_verified"),record.get("source_documents"),record.get("source_locations"),"DRAFT","","",""]
            sheet.append([None]*len(values))
            for index,value in enumerate(values,1):
                if isinstance(value,(list,dict)): value=json.dumps(value,ensure_ascii=False)
                if value is None: value=""
                if isinstance(value,str):
                    value=ILLEGAL_CHARACTERS_RE.sub("",value)
                    if len(value)>32000: value=value[:31900]+" … [full content in review_queue.json]"
                cell=sheet.cell(sheet.max_row,index,value)
                if isinstance(value,str): cell.data_type="s"  # Never execute names/queries as Excel formulas.
        sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="24445C")
        for column in "ABCDEFGHIJKL": sheet.column_dimensions[column].width=28 if column not in "CGH" else 52
    workbook.save(directory/"review_queue.xlsx")
